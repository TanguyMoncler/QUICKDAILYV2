from docx import Document
from pathlib import Path
from datetime import date

from services.market_data_service import get_market_data, get_table_2_data, get_table_3_data
from services.word_service import fill_table_1, fill_table_2, fill_table_3, replace_date, replace_in_doc, clear_remaining_placeholders
import config.settings as settings


if __name__ == "__main__":
    # Crée le dossier output s'il n'existe pas
    output_dir = Path("output")
    output_dir.mkdir(exist_ok=True)

    # Récupération des données marché
    market_data = get_market_data()

    # Chargement des données Table 2 et Table 3 depuis settings
    table_2_data = get_table_2_data(settings.TABLE_2_TICKERS)
    table_3_data = get_table_3_data(settings.SECTOR_TICKERS)

    # Chargement du template Word
    doc = Document("templates/closing_template.docx")

    # Remplissage des tableaux
    fill_table_1(doc, market_data)
    fill_table_2(doc, table_2_data)
    fill_table_3(doc, table_3_data)

    # Remplacement de la date
    replace_date(doc, date.today())

    # Supprime/ remplace les placeholders restants par '-'
    clear_remaining_placeholders(doc)

    # Sauvegarde du document final
    filename = f"{date.today().isoformat()} daily closing.docx"
    doc.save(output_dir / filename)
    print(f"Saved report: {output_dir / filename}")
